"""Write the briefing note as a Word document from the content R wrote (reports/briefing-note.json)."""

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
c = json.loads((ROOT / "reports" / "briefing-note.json").read_text())

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Cm(1.6)
sec.left_margin = sec.right_margin = Cm(2.0)
normal = doc.styles["Normal"]
normal.font.name, normal.font.size = "Arial", Pt(10)
normal.paragraph_format.space_after = Pt(5)
doc.core_properties.author = "Ahmad Bilal Hashimi"

p = doc.add_paragraph("UNCLASSIFIED / NON CLASSIFIÉ")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].bold, p.runs[0].font.size = True, Pt(8.5)
for text, size in (("BRIEFING NOTE", 11), (c["title"], 12.5)):
    r = doc.add_paragraph().add_run(text)
    r.bold, r.font.size = True, Pt(size)

n = 0
for heading, paras in c["sections"]:
    h = doc.add_paragraph()
    hr = h.add_run(heading.upper())
    hr.bold, hr.font.size = True, Pt(10)
    h.paragraph_format.space_before = Pt(8)
    for text in paras:
        n += 1
        doc.add_paragraph(f"{n}. {text}")
    if heading == "Analysis":
        doc.add_picture(str(ROOT / "reports" / "figures" / "fig2_groups.png"), width=Cm(12))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

f = doc.add_paragraph(c["footer"])
f.runs[0].font.size, f.runs[0].italic = Pt(8), True
f.runs[0].font.color.rgb = RGBColor(0x5A, 0x5A, 0x5A)
doc.save(ROOT / "reports" / "briefing-note.docx")
print("reports/briefing-note.docx")
